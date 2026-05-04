from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pathlib import Path

BG = RGBColor(0x1e, 0x1e, 0x2e)
ACCENT = RGBColor(0x89, 0xb4, 0xfa)
TEXT = RGBColor(0xcd, 0xd6, 0xf4)
SUBTEXT = RGBColor(0xa6, 0xe3, 0xa1)
WARN = RGBColor(0xf3, 0x8b, 0xa8)
YELLOW = RGBColor(0xf9, 0xe2, 0xaf)

W, H = Inches(13.33), Inches(7.5)

def new_prs():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    return prs

def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_textbox(slide, text, x, y, w, h, size=24, bold=False, color=TEXT, align=PP_ALIGN.LEFT, italic=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return tb

def add_bullet_slide(prs, title, bullets, notes=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, BG)
    # Title bar
    add_textbox(slide, title, Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.8),
                size=32, bold=True, color=ACCENT)
    # Divider line
    from pptx.util import Pt as pt2
    line = slide.shapes.add_connector(1, Inches(0.5), Inches(1.15), Inches(12.83), Inches(1.15))
    line.line.color.rgb = ACCENT
    line.line.width = Pt(1.5)

    y = Inches(1.3)
    for item in bullets:
        indent = item.get('indent', 0)
        text = item['text']
        color = item.get('color', TEXT)
        size = item.get('size', 22)
        prefix = '  ' * indent + ('• ' if indent == 0 else '  ─ ')
        tb = slide.shapes.add_textbox(Inches(0.6 + indent * 0.3), y, Inches(11.8), Inches(0.55))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = prefix + text
        run.font.size = Pt(size)
        run.font.color.rgb = color
        y += Inches(0.48 + indent * 0.0)
    return slide

prs = new_prs()

# ── Slide 1: Title ──────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BG)
add_textbox(slide, 'Weekly Report', Inches(1), Inches(1.8), Inches(11), Inches(1.2),
            size=52, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
add_textbox(slide, 'Vulkan Persistent Thread Workgraph POC', Inches(1), Inches(3.0), Inches(11), Inches(0.8),
            size=26, color=TEXT, align=PP_ALIGN.CENTER)
add_textbox(slide, '2026 / 04 / 20', Inches(1), Inches(3.8), Inches(11), Inches(0.6),
            size=20, color=SUBTEXT, align=PP_ALIGN.CENTER, italic=True)
add_textbox(slide, 'LUO-WEN-CHIEN', Inches(1), Inches(4.4), Inches(11), Inches(0.5),
            size=18, color=RGBColor(0x6c, 0x70, 0x86), align=PP_ALIGN.CENTER)

# ── Slide 2: 本週做了什麼 (Overview) ────────────────────────────
add_bullet_slide(prs, '本週完成事項 Overview', [
    {'text': '① 科赫雪花（Koch Snowflake）接入 Persistent Thread 管線', 'color': YELLOW, 'size': 24},
    {'text': 'Branching feedback loop：1 條邊 → 4 條子邊，驗證 DAG 任務流', 'indent': 1, 'size': 20},
    {'text': '② Workgroup-Level 角色重構，消除 Wavefront Divergence', 'color': YELLOW, 'size': 24},
    {'text': '角色由 thread-level 改為 workgroup-level，Dispatch 縮為 8 WG', 'indent': 1, 'size': 20},
    {'text': '③ 修復三個跨 CU 記憶體可見性 Bug', 'color': YELLOW, 'size': 24},
    {'text': '最終驗證：精確產生 768 條邊 / 1536 頂點，無垃圾資料', 'indent': 1, 'size': 20, 'color': SUBTEXT},
])

# ── Slide 3: Koch Snowflake ──────────────────────────────────────
add_bullet_slide(prs, '① Koch Snowflake — Branching Feedback Loop', [
    {'text': '任務定義：P1→P2 細分為 4 條子邊（+60° 旋轉），MAX_DEPTH=4', 'color': TEXT},
    {'text': '3 條初始邊 × 4⁴ 遞迴 = 768 條最終邊，1536 個頂點', 'indent': 1, 'size': 20, 'color': SUBTEXT},
    {'text': '三節點拓撲（含 Feedback Loop）', 'color': TEXT},
    {'text': 'Node A  →  Q1  →  Node B', 'indent': 1, 'size': 19, 'color': YELLOW},
    {'text': 'depth < 4 → 推回 Q1（branching × 4）', 'indent': 2, 'size': 18},
    {'text': 'depth == 4 → 轉發 Q2 → Node C → Output Buffer', 'indent': 2, 'size': 18},
    {'text': '資料結構擴充', 'color': TEXT},
    {'text': 'payload[6]：坐標 + depth + per-slot ready flag', 'indent': 1, 'size': 20},
    {'text': '新增第 4 個 SSBO：output vertex buffer（Line List）', 'indent': 1, 'size': 20},
])

# ── Slide 4: Workgroup-Level Role ────────────────────────────────
add_bullet_slide(prs, '② Workgroup-Level 角色重構', [
    {'text': '原設計問題：gId % 32 → 同一 wavefront 走 3 條不同分支', 'color': WARN},
    {'text': 'RDNA3 wave32 = 32 threads，GPU 必須依序執行每條分支並 mask 其他 lane', 'indent': 1, 'size': 20},
    {'text': '重構方案：角色改由 gl_WorkGroupID.x 決定', 'color': SUBTEXT},
    {'text': 'WG 0–5：Node B（Subdivider），WG 0 兼任初始播種', 'indent': 1, 'size': 20},
    {'text': 'WG 6–7：Node C（Writer）', 'indent': 1, 'size': 20},
    {'text': 'Dispatch (32,1,1) → (8,1,1)，同一 workgroup 內所有 thread 做相同工作', 'indent': 1, 'size': 20},
    {'text': '新增 specialization constant NODE_C_START（constant_id=2）', 'color': TEXT},
])

# ── Slide 5: 3 Memory Bugs ──────────────────────────────────────
add_bullet_slide(prs, '③ 三個跨 CU 記憶體可見性 Bug', [
    {'text': 'Bug 1：缺少 coherent → Node C 完全無法消費 Q2', 'color': WARN},
    {'text': '非原子讀取從 L0/L1 讀到 stale 值；修復：所有 SSBO 加 coherent', 'indent': 1, 'size': 19},
    {'text': 'Bug 2：Count 先於 Data → 多 192 條邊 + 垃圾頂點', 'color': WARN},
    {'text': 'Consumer 在 producer 寫完前已讀取 slot；修復：payload[5] per-slot ready flag', 'indent': 1, 'size': 19},
    {'text': 'Bug 3：Plain store 不可靠 → 257 個任務「消失」', 'color': WARN},
    {'text': 'Plain store 停在 write-combining buffer；修復：atomicExchange / atomicCompSwap', 'indent': 1, 'size': 19},
    {'text': '結論：跨 CU 通訊中，coherent + plain store 不夠；所有共享狀態必須原子操作', 'color': YELLOW, 'size': 21},
])

# ── Slide 6: Results & Next Steps ───────────────────────────────
add_bullet_slide(prs, '結果與下週計畫', [
    {'text': '本週驗證結果', 'color': SUBTEXT, 'size': 25},
    {'text': '768 邊 / 1536 頂點精確產生，Q1/Q2 正確歸零', 'indent': 1},
    {'text': 'AMD RX 7900 XTX，耗時 ~2.76s（含 CPU 10ms polling 開銷）', 'indent': 1, 'size': 20},
    {'text': '下週計畫', 'color': SUBTEXT, 'size': 25},
    {'text': '渲染管線：output vertex buffer → VK_PRIMITIVE_TOPOLOGY_LINE_LIST 顯示', 'indent': 1},
    {'text': '性能量測：vkCmdWriteTimestamp 量化 GPU 純計算時間 + CAS 失敗率', 'indent': 1},
    {'text': 'Phase 4 探索：intra-workgroup 通訊移至 LDS（shared memory）', 'indent': 1},
])

out = Path(__file__).resolve().parents[2] / 'reports' / '2026-04' / 'weekly' / 'weekly_report_260420.pptx'
prs.save(str(out))
print(f'Saved: {out}')
