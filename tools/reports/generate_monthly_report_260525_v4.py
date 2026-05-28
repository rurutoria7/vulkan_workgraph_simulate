from __future__ import annotations

import csv
import math
import re
import statistics
from dataclasses import dataclass
from pathlib import Path

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


REPO = Path(__file__).resolve().parents[2]
REPORT = REPO / "reports" / "2026-05"
METRICS = REPORT / "metrics"
MONTHLY = REPORT / "monthly"
DECK_DIR = MONTHLY / "decks"
VERIFY_DIR = MONTHLY / "verification" / "260525_optimization_evaluation_v4"
ASSET_DIR = VERIFY_DIR / "assets"

CSV_OUT = METRICS / "monthly_v4_normalized_measurements_20260525.csv"
NOTES_OUT = METRICS / "monthly_v4_measurement_notes_20260525.md"
PPTX_OUT = DECK_DIR / "260525_MTK_WorkGraph_VulkanPersistentThread_POC_OptimizationEvaluation_v4.pptx"

WIDE_W = Inches(13.333333)
WIDE_H = Inches(7.5)


class C:
    WHITE = "FFFFFF"
    INK = "111111"
    BODY = "333333"
    MUTED = "666666"
    LINE = "D9D9D9"
    LIGHT = "F7F7F7"
    ORANGE = "D97000"
    BLUE = "1F77B4"
    GREEN = "2E7D32"
    GRAY = "6E6E6E"
    RED = "C62828"


SEMANTIC = {
    "problem": C.ORANGE,
    "evidence": C.BLUE,
    "adopted": C.GREEN,
    "neutral": C.GRAY,
    "rejected": C.RED,
}


@dataclass
class MeasurementSpec:
    category: str
    stage: str
    config: str
    source: Path
    warmup_drop: int
    correctness_ok: str = ""
    notes: str = ""


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def rel(path: Path) -> str:
    return path.relative_to(REPO).as_posix()


def read_metric_rows(path: Path) -> list[dict[str, str]]:
    lines = [line.strip() for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    if not lines:
        return []
    if lines[0].startswith("WG_METRICS_HEADER "):
        header = lines[0][len("WG_METRICS_HEADER ") :].split(",")
        data = []
        for line in lines[1:]:
            if line.startswith("WG_METRICS "):
                line = line[len("WG_METRICS ") :]
            data.append(line)
        return list(csv.DictReader(data, fieldnames=header))
    return list(csv.DictReader(lines))


def to_float(value: str) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except ValueError:
        return None


def percentile(values: list[float], pct: float) -> float:
    if not values:
        return math.nan
    ordered = sorted(values)
    if len(ordered) == 1:
        return ordered[0]
    pos = (len(ordered) - 1) * pct
    lo = math.floor(pos)
    hi = math.ceil(pos)
    if lo == hi:
        return ordered[lo]
    weight = pos - lo
    return ordered[lo] * (1.0 - weight) + ordered[hi] * weight


def fmt(value: float | None, digits: int = 4) -> str:
    if value is None or math.isnan(value):
        return ""
    return f"{value:.{digits}f}"


def compute_stats(spec: MeasurementSpec) -> dict[str, str]:
    rows = read_metric_rows(spec.source)
    used = rows[spec.warmup_drop :]
    values = [to_float(row.get("compute_ms", "")) for row in used]
    values = [v for v in values if v is not None]
    if not values:
        return {
            "category": spec.category,
            "stage": spec.stage,
            "config": spec.config,
            "source_file": rel(spec.source),
            "warmup_frames_dropped": str(spec.warmup_drop),
            "sample_count": "0",
            "median_compute_ms": "",
            "avg_compute_ms": "",
            "p10_compute_ms": "",
            "p90_compute_ms": "",
            "min_compute_ms": "",
            "max_compute_ms": "",
            "correctness_ok": spec.correctness_ok,
            "notes": spec.notes,
        }
    return {
        "category": spec.category,
        "stage": spec.stage,
        "config": spec.config,
        "source_file": rel(spec.source),
        "warmup_frames_dropped": str(spec.warmup_drop),
        "sample_count": str(len(values)),
        "median_compute_ms": fmt(statistics.median(values)),
        "avg_compute_ms": fmt(statistics.fmean(values)),
        "p10_compute_ms": fmt(percentile(values, 0.10)),
        "p90_compute_ms": fmt(percentile(values, 0.90)),
        "min_compute_ms": fmt(min(values)),
        "max_compute_ms": fmt(max(values)),
        "correctness_ok": spec.correctness_ok,
        "notes": spec.notes,
    }


def request_batching_rows() -> list[dict[str, str]]:
    summary = METRICS / "request_batching_validation_20260525" / "request_batching_validation_summary.csv"
    rows = list(csv.DictReader(summary.read_text(encoding="utf-8").splitlines()))
    out = []
    for row in rows:
        suite = row["suite"]
        setting = row["setting"]
        case = row["case"]
        raw = summary.parent / f"raw_{suite}_{setting}_{case}_timing.csv"
        notes = "Suite-local request batching comparison."
        if row["timing_samples_used"] in ("0.000000", "0"):
            notes = "No valid timing samples for this run in the validation summary."
        out.append(
            {
                "category": "request_batching",
                "stage": f"{suite}/{setting}",
                "config": case,
                "source_file": rel(raw) if raw.exists() else rel(summary),
                "warmup_frames_dropped": str(int(float(row["warmup_drop_frames_per_run"]))),
                "sample_count": str(int(float(row["timing_samples_used"]))),
                "median_compute_ms": fmt(to_float(row["compute_median_ms"])),
                "avg_compute_ms": fmt(to_float(row["compute_avg_ms"])),
                "p10_compute_ms": fmt(to_float(row["compute_p10_ms"])),
                "p90_compute_ms": fmt(to_float(row["compute_p90_ms"])),
                "min_compute_ms": fmt(to_float(row["compute_min_ms"])),
                "max_compute_ms": "",
                "correctness_ok": row["correctness_ok"],
                "notes": notes,
            }
        )
    return out


def build_measurements() -> list[dict[str, str]]:
    specs = [
        MeasurementSpec(
            "main_timeline",
            "Baseline",
            "Q1 batch on, Q2 scalar, C=24/B=72",
            METRICS / "rgp_occupancy_20260509" / "nodec_24_timestamps.csv",
            1,
            "1",
            "Short historical sweep. Drop first warm frame.",
        ),
        MeasurementSpec(
            "main_timeline",
            "B/C ratio",
            "C=40/B=56",
            METRICS / "rgp_occupancy_20260509" / "nodec_40_timestamps.csv",
            1,
            "1",
            "Short historical sweep. Drop first warm frame.",
        ),
        MeasurementSpec(
            "main_timeline",
            "Q1 sharding",
            "Q1s16/Q2s1, C=38/B=58",
            METRICS / "q1_shards_nodec_20260510" / "q1s_16_nodec_38_refine_timestamps.csv",
            1,
            "1",
            "Drop first warm frame to match the existing summary run.",
        ),
        MeasurementSpec(
            "main_timeline",
            "Q2 sharding",
            "Q1s16/Q2s16, C=16/B=80",
            METRICS / "q2_shards_20260510" / "q1s16_q2s16_nodec16_timestamps.csv",
            2,
            "1",
            "Two warm frames are dropped to match the existing Q2 sharding summary.",
        ),
        MeasurementSpec(
            "main_timeline",
            "Q1 lane-pop",
            "Q1s16/Q2s16, C=24/B=72",
            METRICS / "q2_shards_20260510" / "q1s16_q2s16_lanepop_c24_timestamps.csv",
            30,
            "1",
            "Longer timing run. Drop 30 warm frames.",
        ),
        MeasurementSpec(
            "main_timeline",
            "High shard scaling",
            "Q1s256/Q2s256, C=24/B=72",
            METRICS / "q2_shards_20260510" / "q1s256_q2s256_lanepop_c24_max256_timestamps.csv",
            30,
            "1",
            "Longer timing run. Drop 30 warm frames.",
        ),
        MeasurementSpec(
            "final_confirmation",
            "Final repeat 1",
            "Main integrated defaults",
            METRICS / "q2_shards_20260510" / "main_integrated_repeat_1.csv",
            30,
            "1",
            "Main worktree confirmation repeat.",
        ),
        MeasurementSpec(
            "final_confirmation",
            "Final repeat 2",
            "Main integrated defaults",
            METRICS / "q2_shards_20260510" / "main_integrated_repeat_2.csv",
            30,
            "1",
            "Main worktree confirmation repeat.",
        ),
        MeasurementSpec(
            "final_confirmation",
            "Final repeat 3",
            "Main integrated defaults",
            METRICS / "q2_shards_20260510" / "main_integrated_repeat_3.csv",
            30,
            "1",
            "Main worktree confirmation repeat.",
        ),
        MeasurementSpec(
            "final_confirmation",
            "Final repeat 4",
            "Main integrated defaults",
            METRICS / "q2_shards_20260510" / "main_integrated_repeat_4.csv",
            30,
            "1",
            "Main worktree confirmation repeat.",
        ),
    ]

    rows = [compute_stats(spec) for spec in specs]

    for shards, c, source in [
        (16, 24, "q1s16_q2s16_lanepop_c24_timestamps.csv"),
        (32, 28, "q1s32_q2s32_lanepop_c28_timestamps.csv"),
        (64, 28, "q1s64_q2s64_lanepop_c28_timestamps.csv"),
        (128, 24, "q1s128_q2s128_lanepop_c24_max128_timestamps.csv"),
        (192, 24, "q1s192_q2s192_lanepop_c24_max256_timestamps.csv"),
        (256, 24, "q1s256_q2s256_lanepop_c24_max256_timestamps.csv"),
        (384, 24, "q1s384_q2s384_lanepop_c24_max512_timestamps.csv"),
        (512, 28, "q1s512_q2s512_lanepop_c28_max512_timestamps.csv"),
    ]:
        rows.append(
            compute_stats(
                MeasurementSpec(
                    "shard_scaling",
                    f"Q1/Q2 shards {shards}",
                    f"Q1s{shards}/Q2s{shards}, C={c}/B={96-c}",
                    METRICS / "q2_shards_20260510" / source,
                    30,
                    "1",
                    "Long shard scaling sweep. Drop 30 warm frames.",
                )
            )
        )

    rows.extend(request_batching_rows())
    return rows


FIELDS = [
    "category",
    "stage",
    "config",
    "source_file",
    "warmup_frames_dropped",
    "sample_count",
    "median_compute_ms",
    "avg_compute_ms",
    "p10_compute_ms",
    "p90_compute_ms",
    "min_compute_ms",
    "max_compute_ms",
    "correctness_ok",
    "notes",
]


def write_measurement_files(rows: list[dict[str, str]]) -> None:
    METRICS.mkdir(parents=True, exist_ok=True)
    with CSV_OUT.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)

    timeline = [row for row in rows if row["category"] == "main_timeline"]
    final = [row for row in rows if row["category"] == "final_confirmation"]
    request = [row for row in rows if row["category"] == "request_batching"]

    def md_table(rows_: list[dict[str, str]], cols: list[str]) -> str:
        header = "| " + " | ".join(cols) + " |"
        sep = "| " + " | ".join(["---"] * len(cols)) + " |"
        body = []
        for row in rows_:
            body.append("| " + " | ".join(row.get(col, "") for col in cols) + " |")
        return "\n".join([header, sep, *body])

    notes = f"""# Monthly v4 Measurement Notes

Date: 2026-05-25

Measurement rules used by the v4 monthly deck.

## Rules

- Headline performance numbers use median clean `compute_ms` after warmup.
- Average remains secondary context; median sets headline speedup.
- Final confirmation reports median, P10, P90, and min from clean timestamp runs.
- Counter-enabled runs are structural evidence only because shader counters add atomic traffic.
- RGP occupancy supplies auxiliary context. Treat `wavefront = 0` as context only; clean timestamps carry dispatch-duration proof.
- Warmup policy is source-specific for historical runs and is recorded in the CSV instead of being hidden.
- Request batching is compared only against the suite-local `q2_batch_off` row.

## Normalized Timeline

{md_table(timeline, ["stage", "config", "warmup_frames_dropped", "sample_count", "median_compute_ms", "avg_compute_ms"])}

## Final Confirmation

{md_table(final, ["stage", "sample_count", "median_compute_ms", "p10_compute_ms", "p90_compute_ms", "min_compute_ms"])}

## Request Batching Validation

{md_table(request, ["stage", "config", "sample_count", "median_compute_ms", "p90_compute_ms", "correctness_ok"])}

## Source Of Truth

Numeric source: `{rel(CSV_OUT)}` plus measurement notes. The CSV contains the source file and warmup rule for each row.
"""
    NOTES_OUT.write_text(notes, encoding="utf-8")


def row_by(rows: list[dict[str, str]], category: str, stage: str) -> dict[str, str]:
    for row in rows:
        if row["category"] == category and row["stage"] == stage:
            return row
    raise KeyError((category, stage))


def request_row(rows: list[dict[str, str]], stage_suffix: str, config: str) -> dict[str, str]:
    for row in rows:
        if row["category"] == "request_batching" and row["stage"].endswith(stage_suffix) and row["config"] == config:
            return row
    raise KeyError((stage_suffix, config))


def create_chart(path: Path, title: str, labels: list[str], values: list[float], color: str = C.BLUE) -> None:
    fig, ax = plt.subplots(figsize=(8.4, 2.9), dpi=180)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    ax.plot(labels, values, color=f"#{color}", linewidth=2.5, marker="o", markersize=5)
    ax.set_title(title, loc="left", fontsize=13, color="#111111", fontweight="bold")
    ax.set_ylabel("Median compute_ms", fontsize=10)
    ax.grid(True, axis="y", color="#DDDDDD", linewidth=0.7)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color("#BBBBBB")
    ax.spines["bottom"].set_color("#BBBBBB")
    ax.tick_params(axis="both", labelsize=9, colors="#333333")
    for i, v in enumerate(values):
        ax.annotate(f"{v:.2f}", (i, v), xytext=(0, 8), textcoords="offset points", ha="center", fontsize=8.5)
    fig.tight_layout(pad=1.0)
    fig.savefig(path, facecolor="white", bbox_inches="tight")
    plt.close(fig)


def create_assets(rows: list[dict[str, str]]) -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    timeline_rows = [row for row in rows if row["category"] == "main_timeline"]
    create_chart(
        ASSET_DIR / "timeline_median_compute.png",
        "Median clean compute_ms after each queue change",
        [row["stage"] for row in timeline_rows],
        [float(row["median_compute_ms"]) for row in timeline_rows],
        C.GREEN,
    )
    shard_rows = [row for row in rows if row["category"] == "shard_scaling"]
    shard_rows.sort(key=lambda row: int(re.search(r"(\d+)$", row["stage"]).group(1)))
    create_chart(
        ASSET_DIR / "high_shard_scaling_median.png",
        "Equal Q1/Q2 shard sweep",
        [row["stage"].split()[-1] for row in shard_rows],
        [float(row["median_compute_ms"]) for row in shard_rows],
        C.BLUE,
    )
    return {
        "timeline": ASSET_DIR / "timeline_median_compute.png",
        "shards": ASSET_DIR / "high_shard_scaling_median.png",
    }


def blank_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(C.WHITE)
    return slide


def textbox(slide, x, y, w, h, text, size=14, color=C.BODY, bold=False, align="left", margin=0.05):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    p.font.name = "Arial"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = rgb(color)
    return shape


def multi_text(slide, x, y, w, h, lines, size=13, color=C.BODY, bullet=False, line_spacing=1.0):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"- {line}" if bullet else line
        p.font.name = "Arial"
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.space_after = Pt(5 * line_spacing)
    return shape


def rect(slide, x, y, w, h, fill=C.WHITE, line=C.LINE, width=1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(width)
    try:
        shape.shadow.inherit = False
    except Exception:
        pass
    return shape


def semantic_tag(slide, x, y, text, kind, w=1.25):
    shape = rect(slide, x, y, w, 0.28, fill=SEMANTIC[kind], line=SEMANTIC[kind], width=0)
    textbox(slide, x, y + 0.03, w, 0.18, text, size=8.5, color=C.WHITE, bold=True, align="center", margin=0)
    return shape


def metric_box(slide, x, y, w, h, value, label, kind="evidence", sub=""):
    rect(slide, x, y, w, h, fill=C.WHITE, line=C.LINE)
    rect(slide, x, y, w, 0.08, fill=SEMANTIC[kind], line=SEMANTIC[kind], width=0)
    textbox(slide, x + 0.08, y + 0.18, w - 0.16, 0.42, value, size=25, color=SEMANTIC[kind], bold=True, align="center")
    textbox(slide, x + 0.1, y + 0.66, w - 0.2, 0.32, label, size=10.5, color=C.BODY, bold=True, align="center")
    if sub:
        textbox(slide, x + 0.1, y + 0.98, w - 0.2, 0.25, sub, size=8.5, color=C.MUTED, align="center")


def header(slide, title: str, page: int):
    textbox(slide, 0.45, 0.24, 8.2, 0.18, "Vulkan Persistent Thread WorkGraph POC", size=8.5, color=C.MUTED, bold=True)
    textbox(slide, 0.45, 0.43, 10.8, 0.52, title, size=26, color=C.INK, bold=True)
    semantic_tag(slide, 12.25, 0.31, f"{page:02d}", "evidence", w=0.42)
    slide.shapes.add_connector(1, Inches(0.45), Inches(0.98), Inches(12.88), Inches(0.98)).line.color.rgb = rgb(C.LINE)
    textbox(slide, 0.45, 7.12, 3.2, 0.18, "2026-05-25", size=7, color=C.MUTED)
    textbox(slide, 9.2, 7.12, 3.3, 0.18, "Optimization Evaluation v4", size=7, color=C.MUTED, align="right")


def add_table(slide, x, y, w, h, columns, rows, col_widths=None, font_size=8.5, header_fill=C.LIGHT, header_color=C.INK):
    table_shape = slide.shapes.add_table(len(rows) + 1, len(columns), Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    if col_widths:
        for idx, width in enumerate(col_widths):
            table.columns[idx].width = Inches(width)
    for cidx, col in enumerate(columns):
        cell = table.cell(0, cidx)
        cell.text = col
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(header_fill)
        for p in cell.text_frame.paragraphs:
            p.font.name = "Arial"
            p.font.size = Pt(font_size)
            p.font.bold = True
            p.font.color.rgb = rgb(header_color)
    for ridx, row in enumerate(rows, 1):
        for cidx, value in enumerate(row):
            cell = table.cell(ridx, cidx)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(C.WHITE)
            for p in cell.text_frame.paragraphs:
                p.font.name = "Arial"
                p.font.size = Pt(font_size)
                p.font.color.rgb = rgb(C.BODY)
    for row in table.rows:
        for cell in row.cells:
            cell.margin_left = Inches(0.04)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
    return table_shape


def arrow(slide, x1, y1, x2, y2, color=C.GRAY):
    conn = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = rgb(color)
    conn.line.width = Pt(1.6)
    conn.line.end_arrowhead = True
    return conn


def box_label(slide, x, y, w, h, title, subtitle="", kind="neutral"):
    rect(slide, x, y, w, h, fill=C.WHITE, line=SEMANTIC[kind], width=1.4)
    textbox(slide, x + 0.08, y + 0.12, w - 0.16, 0.25, title, size=11.5, color=SEMANTIC[kind], bold=True, align="center")
    if subtitle:
        textbox(slide, x + 0.08, y + 0.45, w - 0.16, h - 0.5, subtitle, size=8.7, color=C.BODY, align="center")


def stage_value(rows: list[dict[str, str]], stage: str) -> float:
    return float(row_by(rows, "main_timeline", stage)["median_compute_ms"])


def build_deck(rows: list[dict[str, str]], assets: dict[str, Path]) -> None:
    prs = Presentation()
    prs.slide_width = WIDE_W
    prs.slide_height = WIDE_H
    page = 1

    baseline = stage_value(rows, "Baseline")
    final = float(row_by(rows, "final_confirmation", "Final repeat 1")["median_compute_ms"])
    speedup = baseline / final

    slide = blank_slide(prs)
    textbox(slide, 0.72, 0.45, 3.1, 0.23, "MONTHLY REPORT", size=10, color=C.BLUE, bold=True)
    textbox(slide, 0.72, 0.95, 9.3, 1.08, "Vulkan Persistent Thread\nWorkGraph POC", size=42, color=C.INK, bold=True)
    textbox(slide, 0.75, 2.35, 10.9, 0.52, "May 2026 evaluation: queue atomics set the bottleneck, and sharding removes it", size=18, color=C.BODY, bold=True)
    metric_box(slide, 0.75, 3.25, 2.45, 1.35, f"{baseline:.2f} ms", "Starting median", "problem", "clean timestamp")
    metric_box(slide, 3.55, 3.25, 2.45, 1.35, f"{final:.2f} ms", "Final repeat median", "adopted", "main defaults")
    metric_box(slide, 6.35, 3.25, 2.45, 1.35, f"{speedup:.1f}x", "Median speedup", "evidence", "same workload")
    textbox(slide, 0.75, 5.1, 10.85, 0.72, "Adopt Q1/Q2 sharding and lane-based Q1 pop. Keep request batching out of the main path because the default 256-shard suite regresses.", size=15, color=C.BODY)
    textbox(slide, 0.75, 6.85, 3.2, 0.2, "AMD RX 7900 XTX / RDNA 3", size=8, color=C.MUTED)
    page += 1

    slide = blank_slide(prs)
    header(slide, "Queue atomics explain the measured gains", page)
    cols = [
        ("Bottleneck", "Clean median falls when a change reduces shared Q1/Q2 queue pressure.", "problem"),
        ("Boundary", "Timestamp-only runs carry duration. Counter runs explain CAS, backlog, and ready-spin.", "evidence"),
        ("Adopt", "Q1 shards, Q2 shards, and lane-based Q1 pop reach about 0.99 ms median.", "adopted"),
        ("Reject", "Q2 request batching pushes the default 256-shard suite into the 11 ms band.", "rejected"),
    ]
    for i, (title, body, kind) in enumerate(cols):
        x = 0.65 + i * 3.05
        rect(slide, x, 1.45, 2.7, 3.35, fill=C.WHITE, line=C.LINE)
        semantic_tag(slide, x + 0.18, 1.66, title.upper(), kind, w=1.25)
        textbox(slide, x + 0.2, 2.15, 2.3, 0.38, title, size=16, color=SEMANTIC[kind], bold=True)
        textbox(slide, x + 0.2, 2.72, 2.3, 1.6, body, size=13, color=C.BODY)
    textbox(slide, 0.7, 5.35, 11.7, 0.58, f"The adopted path moves {baseline:.2f} ms to {final:.2f} ms without using shader counters in the timing run.", size=13.5, color=C.INK, bold=True)
    page += 1

    slide = blank_slide(prs)
    header(slide, "Only Q1/Q2 handoff points are shared global queue state", page)
    y = 2.0
    box_label(slide, 0.65, y, 1.35, 0.85, "Node A", "Seed edges once", "neutral")
    box_label(slide, 2.45, y, 1.65, 0.85, "Q1", "atomic count/head/tail\nready flag", "problem")
    box_label(slide, 4.55, y, 1.55, 0.85, "Node B", "Subdivide work", "neutral")
    box_label(slide, 6.55, y, 1.65, 0.85, "Q2", "atomic count/head/tail\nready flag", "problem")
    box_label(slide, 8.65, y, 1.55, 0.85, "Node C", "Write vertices", "neutral")
    box_label(slide, 10.65, y, 1.65, 0.85, "VB Render", "Deterministic draw", "neutral")
    for x1, x2 in [(2.0, 2.45), (4.1, 4.55), (6.1, 6.55), (8.2, 8.65), (10.2, 10.65)]:
        arrow(slide, x1, y + 0.42, x2, y + 0.42, C.GRAY)
    rect(slide, 0.65, 3.65, 5.55, 1.55, fill=C.WHITE, line=C.LINE)
    textbox(slide, 0.9, 3.85, 4.9, 0.32, "Shared fields", size=15.5, color=C.ORANGE, bold=True)
    multi_text(slide, 0.9, 4.28, 4.9, 0.8, ["Q1/Q2 count, head, tail, and ready flags synchronize cross-workgroup handoff.", "Many lanes and workgroups touch the same fields before sharding."], size=12, bullet=True)
    rect(slide, 6.75, 3.65, 5.55, 1.55, fill=C.WHITE, line=C.LINE)
    textbox(slide, 7.0, 3.85, 4.9, 0.32, "Test question", size=15.5, color=C.BLUE, bold=True)
    multi_text(slide, 7.0, 4.28, 4.9, 0.8, ["Candidate changes must reduce clean duration, queue contention, or both.", "Reject cost shifts into polling or ready-spin."], size=12, bullet=True)
    page += 1

    slide = blank_slide(prs)
    header(slide, "Clean timing plus counters define the adoption rule", page)
    rect(slide, 0.7, 1.45, 5.55, 4.25, fill=C.WHITE, line=C.LINE)
    semantic_tag(slide, 0.95, 1.72, "PRIOR REPORT", "neutral", w=1.25)
    textbox(slide, 0.95, 2.18, 4.9, 0.42, "Before May 10", size=18, color=C.INK, bold=True)
    multi_text(slide, 0.95, 2.75, 4.8, 1.55, ["The POC produced correct deterministic output.", "Persistent workers exchanged work through global atomic queues.", "Duration runs still carried diagnostic counters."], size=13.5, bullet=True)
    rect(slide, 7.0, 1.45, 5.55, 4.25, fill=C.WHITE, line=C.LINE)
    semantic_tag(slide, 7.25, 1.72, "RULE", "evidence", w=1.25)
    textbox(slide, 7.25, 2.18, 4.9, 0.42, "Decision rule", size=18, color=C.BLUE, bold=True)
    multi_text(slide, 7.25, 2.75, 4.8, 1.55, ["Clean timestamp runs set the performance number.", "Counter-enabled runs diagnose CAS retries, backlog, and polling.", "Adopt clean-timing improvements."], size=13.5, bullet=True)
    textbox(slide, 0.85, 6.1, 11.3, 0.55, "The conclusion rests on a narrow rule: duration proves impact; counters explain the mechanism.", size=14.5, color=C.INK, bold=True)
    page += 1

    slide = blank_slide(prs)
    header(slide, "Timing and counters answer different questions", page)
    rows2 = [
        ["Clean timestamp runs", "Performance claim", "Median compute_ms after warmup", "Duration, speedup, adoption"],
        ["Counter-enabled runs", "Mechanism check", "CAS fail ratio, backlog, ready-spin, empty-pop", "Queue contention diagnosis"],
        ["RGP occupancy", "Context", "Wavefront view and event timing", "Profiler boundary"],
    ]
    add_table(slide, 0.75, 1.45, 11.85, 2.25, ["Run type", "Purpose", "Primary readout", "What it supports"], rows2, [2.2, 2.1, 3.6, 3.95], 10.5)
    metric_box(slide, 0.85, 4.25, 3.6, 1.15, "Adopt", "Clean runs decide if a change stays.", "evidence")
    metric_box(slide, 4.85, 4.25, 3.6, 1.15, "Diagnose", "Counters identify the queue failure mode.", "problem")
    metric_box(slide, 8.85, 4.25, 3.6, 1.15, "Bound", "RGP views set scope; timing carries proof.", "neutral")
    textbox(slide, 0.9, 5.95, 11.3, 0.55, "Separating run types keeps counter instrumentation out of the performance result.", size=13.5, color=C.INK, bold=True)
    page += 1

    slide = blank_slide(prs)
    header(slide, "Each atomic metric maps to a queue failure mode", page)
    metric_rows = [
        ["CAS fail per success", "Lanes or workgroups retry the same queue state.", "High contention; sharding should reduce the retry rate."],
        ["Attempts vs success", "The control path burns operations without throughput.", "Detects wasted polling."],
        ["Queue high-water", "Backlog exposes producer/consumer imbalance.", "Guides B/C balance and Q2 pressure checks."],
        ["Ready-spin / empty-pop", "Consumers claim or probe before useful work is ready.", "Explains request batching regressions."],
        ["Clean vs counter runs", "The run type separates duration from diagnosis.", "Keeps cause analysis out of headline timing."],
    ]
    add_table(slide, 0.65, 1.35, 12.05, 3.6, ["Metric", "Interpretation", "Decision use"], metric_rows, [2.45, 4.4, 5.2], 9.8)
    textbox(slide, 0.8, 5.35, 11.6, 0.65, "A change earns confidence when clean duration improves and the relevant counter pressure moves in the same direction.", size=14.5, color=C.INK, bold=True)
    page += 1

    slide = blank_slide(prs)
    header(slide, "Queue-locality changes drove a 31.0x median speedup", page)
    timeline = [row for row in rows if row["category"] == "main_timeline"]
    slide.shapes.add_picture(str(assets["timeline"]), Inches(0.7), Inches(1.25), width=Inches(6.25))
    metric_box(slide, 7.35, 1.33, 2.35, 1.15, f"{baseline:.2f} ms", "Starting median", "problem", "clean timestamp")
    metric_box(slide, 10.0, 1.33, 2.35, 1.15, f"{final:.2f} ms", "Final repeat", "adopted", "main defaults")
    metric_box(slide, 8.68, 2.82, 2.35, 1.15, f"{speedup:.1f}x", "Median speedup", "evidence", "same workload")
    stage_kinds = ["problem", "evidence", "adopted", "adopted", "adopted", "adopted"]
    for idx, row in enumerate(timeline):
        x = 0.72 + idx * 2.05
        kind = stage_kinds[idx]
        rect(slide, x, 4.85, 1.82, 0.84, fill=C.WHITE, line=SEMANTIC[kind], width=1.25)
        textbox(slide, x + 0.08, 4.98, 1.66, 0.2, row["stage"], size=8.3, color=SEMANTIC[kind], bold=True, align="center", margin=0)
        textbox(slide, x + 0.08, 5.25, 1.66, 0.22, f"{row['median_compute_ms']} ms", size=8.2, color=C.INK, bold=True, align="center", margin=0)
    textbox(slide, 0.75, 6.05, 11.3, 0.55, "Every large drop follows a queue split or scan-diversification step; request batching stays outside the adopted path.", size=14.5, color=C.GREEN, bold=True)
    page += 1

    slide = blank_slide(prs)
    header(slide, "C/B tuning reduces backlog; sharding remains necessary", page)
    bc_rows = []
    for c in [24, 32, 40, 48, 64]:
        file = METRICS / "rgp_occupancy_20260509" / f"nodec_{c}_timestamps.csv"
        spec = MeasurementSpec("bc", "B/C", f"C={c}/B={96-c}", file, 1)
        stat = compute_stats(spec)
        bc_rows.append([f"{c}/{96-c}", stat["median_compute_ms"], stat["avg_compute_ms"]])
    add_table(slide, 0.7, 1.45, 4.3, 2.45, ["C/B", "Median ms", "Avg ms"], bc_rows, [1.2, 1.45, 1.45], 10.5)
    rect(slide, 5.55, 1.45, 6.75, 2.25, fill=C.WHITE, line=C.LINE)
    semantic_tag(slide, 5.8, 1.7, "READOUT", "evidence", w=1.0)
    textbox(slide, 5.8, 2.15, 5.95, 0.36, "Writer balance is a setup fix", size=17, color=C.BLUE, bold=True)
    multi_text(slide, 5.8, 2.65, 5.95, 0.9, ["C=40/B=56 lowers Q2 high-water from about 50k to about 9k.", "Median improves by 2.86 ms, but the queue remains a shared atomic hotspot."], size=12.5, bullet=True)
    metric_box(slide, 0.9, 4.55, 2.55, 1.1, "50k -> 9k", "Q2 high-water trend", "evidence", "counter-enabled runs")
    metric_box(slide, 3.8, 4.55, 2.55, 1.1, "30.75 -> 27.89", "Median ms", "adopted", "clean runs")
    textbox(slide, 6.85, 4.72, 5.1, 0.8, "Set C=40/B=56 before testing structural queue changes.", size=13.5, color=C.INK, bold=True)
    page += 1

    slide = blank_slide(prs)
    header(slide, "Q1 sharding cuts median to 13.74 ms and exposes Q2", page)
    q1_summary = list(csv.DictReader((METRICS / "q1_shards_nodec_20260510" / "q1_shards_nodec_summary.csv").read_text(encoding="utf-8").splitlines()))
    q1_refine = list(csv.DictReader((METRICS / "q1_shards_nodec_20260510" / "q1s16_nodec_refine_summary.csv").read_text(encoding="utf-8").splitlines()))
    wanted = [("4", "40", "56"), ("8", "40", "56"), ("16", "36", "60"), ("16", "38", "58"), ("16", "40", "56")]
    q1_rows = []
    for q, c, b in wanted:
        if q == "16" and c in {"36", "38", "40"}:
            match = next(r for r in q1_refine if r["c_workgroups"] == c and r["b_workgroups"] == b)
        else:
            match = next(r for r in q1_summary if r["q1_shards"] == q and r["c_workgroups"] == c and r["b_workgroups"] == b)
        q1_rows.append([q, f"{c}/{b}", match["median_ms"], match["avg_ms"]])
    add_table(slide, 0.65, 1.35, 5.2, 2.85, ["Q1 shards", "C/B", "Median ms", "Avg ms"], q1_rows, [1.25, 1.05, 1.35, 1.35], 9.6)
    box_label(slide, 6.5, 1.55, 1.9, 0.95, "Before", "Single Q1 queue\nhot count/head/tail", "problem")
    arrow(slide, 8.55, 2.02, 9.25, 2.02, C.GRAY)
    box_label(slide, 9.45, 1.55, 2.4, 0.95, "After", "Q1 split across shards\nsame queue protocol", "adopted")
    multi_text(slide, 6.55, 3.0, 5.5, 1.35, ["Q1 sharding turns one global atomic queue into 16 smaller queues.", "After Q1 pressure drops, Q2 count/head/tail becomes the next bottleneck target."], size=12.3, bullet=True)
    metric_box(slide, 0.9, 4.75, 2.65, 1.05, "13.74 ms", "Q1s16/Q2s1 median", "adopted")
    textbox(slide, 4.0, 4.92, 7.8, 0.65, "The result stacks with B/C tuning, so the gain tracks queue contention instead of an unrelated shader change.", size=14, color=C.INK, bold=True)
    page += 1

    slide = blank_slide(prs)
    header(slide, "Q2 sharding removes the second global queue wall", page)
    q2_summary = list(csv.DictReader((METRICS / "q2_shards_20260510" / "q2_shards_summary.csv").read_text(encoding="utf-8").splitlines()))
    q2_rows = [[r["q2_shards"], r["median_ms"], r["avg_ms"], r["speedup_vs_q2s1"]] for r in q2_summary]
    add_table(slide, 0.65, 1.35, 4.9, 2.85, ["Q2 shards", "Median ms", "Avg ms", "Speedup"], q2_rows, [1.15, 1.25, 1.25, 1.15], 9.6)
    rect(slide, 6.0, 1.35, 6.2, 2.65, fill=C.WHITE, line=C.LINE)
    semantic_tag(slide, 6.25, 1.63, "DECISION", "adopted", w=1.0)
    textbox(slide, 6.25, 2.08, 5.5, 0.38, "Shard Q2 before batching it", size=17, color=C.GREEN, bold=True)
    multi_text(slide, 6.25, 2.6, 5.5, 1.0, ["Q2 median improves from 1 to 16 shards.", "After B/C rebalance, Q1s16/Q2s16 reaches about 5.57 ms median."], size=12.3, bullet=True)
    metric_box(slide, 0.9, 4.75, 2.65, 1.05, "5.57 ms", "Q1s16/Q2s16 median", "adopted")
    metric_box(slide, 3.9, 4.75, 2.65, 1.05, "2.79x", "vs Q2s1 in sweep", "evidence")
    textbox(slide, 6.95, 4.9, 5.1, 0.65, "The same pattern repeats: split the shared queue state, then retune worker balance.", size=14, color=C.INK, bold=True)
    page += 1

    slide = blank_slide(prs)
    header(slide, "Lane-pop removes same-shard starts inside a workgroup", page)
    lane_rows = [
        ["lane-pop off, C=16/B=80", "5.60", "6.01", "5.54"],
        ["lane-pop on, C=16/B=80", "4.00", "4.39", "3.90"],
        ["lane-pop on, C=24/B=72", row_by(rows, "main_timeline", "Q1 lane-pop")["median_compute_ms"], row_by(rows, "main_timeline", "Q1 lane-pop")["avg_compute_ms"], row_by(rows, "main_timeline", "Q1 lane-pop")["min_compute_ms"]],
    ]
    add_table(slide, 0.65, 1.35, 6.1, 2.35, ["Mode", "Median ms", "Avg ms", "Min ms"], lane_rows, [3.25, 1.0, 0.9, 0.85], 9.4)
    rect(slide, 7.2, 1.35, 5.15, 2.15, fill=C.WHITE, line=C.LINE)
    textbox(slide, 7.45, 1.65, 4.55, 0.36, "Scan-start rule", size=17, color=C.BLUE, bold=True)
    textbox(slide, 7.45, 2.1, 4.45, 0.45, "Old: wgId % Q1_QUEUE_SHARDS", size=12, color=C.ORANGE, bold=True)
    textbox(slide, 7.45, 2.62, 4.45, 0.5, "New: (wgId * local_size_x + localId) % Q1_QUEUE_SHARDS", size=11, color=C.GREEN, bold=True)
    textbox(slide, 0.85, 4.35, 11.4, 0.9, "Lane-pop keeps the queue protocol intact. It changes where each lane begins scanning, so the improvement isolates shard-selection pressure.", size=14.5, color=C.INK, bold=True)
    page += 1

    slide = blank_slide(prs)
    header(slide, "256/256 is the measured knee of the shard-count tradeoff", page)
    slide.shapes.add_picture(str(assets["shards"]), Inches(0.7), Inches(1.35), width=Inches(7.0))
    shard_rows = [row for row in rows if row["category"] == "shard_scaling"]
    shard_rows.sort(key=lambda row: int(row["stage"].split()[-1]))
    compact = [[row["stage"].split()[-1], row["median_compute_ms"], row["p90_compute_ms"]] for row in shard_rows]
    add_table(slide, 8.15, 1.35, 4.1, 3.3, ["Q1/Q2", "Median", "P90"], compact, [1.15, 1.4, 1.4], 9.0)
    metric_box(slide, 0.95, 5.35, 2.6, 1.05, "256 / 256", "Adopted default", "adopted")
    textbox(slide, 4.0, 5.45, 7.9, 0.75, "The 256-shard point has the best median in the sweep; 384 and 512 add probe/index cost without improving the tail.", size=14, color=C.INK, bold=True)
    page += 1

    slide = blank_slide(prs)
    header(slide, "Request batching targets dequeue CAS first", page)
    box_label(slide, 0.75, 1.55, 2.25, 1.0, "Problem", "Per-item dequeue CAS\ncreates heavy retry traffic", "problem")
    arrow(slide, 3.2, 2.05, 4.0, 2.05, C.GRAY)
    box_label(slide, 4.15, 1.55, 2.25, 1.0, "Idea", "One subgroup claim\nreserves multiple slots", "evidence")
    arrow(slide, 6.6, 2.05, 7.4, 2.05, C.GRAY)
    box_label(slide, 7.55, 1.55, 2.25, 1.0, "Expected", "Fewer CAS attempts\nper consumed item", "adopted")
    arrow(slide, 10.0, 2.05, 10.8, 2.05, C.GRAY)
    box_label(slide, 10.95, 1.55, 1.65, 1.0, "Risk", "More ready-spin\nor empty slots", "problem")
    rect(slide, 0.8, 3.55, 11.7, 1.45, fill=C.WHITE, line=C.LINE)
    textbox(slide, 1.05, 3.85, 11.0, 0.38, "Adoption test", size=17, color=C.BLUE, bold=True)
    textbox(slide, 1.05, 4.32, 10.9, 0.5, "A batch claim must lower clean median and P90 in the default 256-shard suite without moving cost into ready-spin or empty-slot work.", size=13.5, color=C.INK, bold=True)
    page += 1

    slide = blank_slide(prs)
    header(slide, "Default 256-shard batching regresses, so it stays out", page)
    default_off = request_row(rows, "default_256_shards", "q2_batch_off")
    metric_box(slide, 0.75, 1.25, 2.9, 1.15, f"{default_off['median_compute_ms']} ms", "Default off baseline", "neutral", "suite-local")
    metric_box(slide, 4.0, 1.25, 2.9, 1.15, "11 ms range", "Default batching result", "rejected", "L8/L16/L32")
    textbox(slide, 7.45, 1.32, 4.7, 0.95, "Stress cases can improve, but the default path moves from 1.94 ms off-baseline into the 11 ms band.", size=14, color=C.RED, bold=True)
    default_rows = []
    for case in ["q2_batch_off", "q2_batch_l8", "q2_batch_l16", "q2_batch_l32", "q2_batch_l1", "q2_batch_l2", "q2_batch_l4"]:
        r = request_row(rows, "default_256_shards", case)
        default_rows.append([case.replace("q2_batch_", ""), r["sample_count"], r["median_compute_ms"] or "no samples", r["p90_compute_ms"] or "no samples"])
    add_table(slide, 0.6, 3.05, 6.0, 2.45, ["Default 256 suite", "Samples", "Median", "P90"], default_rows, [1.7, 1.05, 1.5, 1.5], 8.3)
    stress_rows = []
    for case in ["q2_batch_off", "q2_batch_l4", "q2_batch_l8", "q2_batch_l16", "q2_batch_l32"]:
        r = request_row(rows, "stress_16_shards", case)
        stress_rows.append([case.replace("q2_batch_", ""), r["median_compute_ms"], r["p90_compute_ms"]])
    add_table(slide, 6.95, 3.05, 5.6, 2.1, ["Stress 16 suite", "Median", "P90"], stress_rows, [2.0, 1.7, 1.7], 8.4)
    textbox(slide, 0.82, 5.9, 11.1, 0.45, "Decision: reject batching for the current main protocol; revisit it only after the queue protocol changes.", size=14.5, color=C.RED, bold=True)
    page += 1

    slide = blank_slide(prs)
    header(slide, "The main path encodes the proven queue changes", page)
    metric_box(slide, 0.85, 1.35, 2.6, 1.2, "256", "Q1 shards", "adopted")
    metric_box(slide, 3.75, 1.35, 2.6, 1.2, "256", "Q2 shards", "adopted")
    metric_box(slide, 6.65, 1.35, 2.6, 1.2, "ON", "Q1 lane-pop", "adopted")
    metric_box(slide, 9.55, 1.35, 2.6, 1.2, "24 / 72", "C/B workers", "adopted")
    rows3 = [
        ["MAX_QUEUE_SHARDS", "256"],
        ["Q1_QUEUE_SHARDS_DEFAULT", "256"],
        ["Q2_QUEUE_SHARDS_DEFAULT", "256"],
        ["ENABLE_Q1_LANE_POP", "default on"],
        ["NODE_C_START", "72, so C=24/B=72"],
    ]
    add_table(slide, 0.9, 3.25, 5.4, 2.25, ["Setting", "Value"], rows3, [3.4, 1.9], 10.2)
    textbox(slide, 7.0, 3.45, 5.1, 0.9, "The defaults preserve the measured mechanisms: split both queues, diversify Q1 scan starts, and keep the B/C mix from the best clean run.", size=14.5, color=C.INK, bold=True)
    page += 1

    slide = blank_slide(prs)
    header(slide, "Four clean repeats keep median and P90 near 1 ms", page)
    final_rows = [row for row in rows if row["category"] == "final_confirmation"]
    add_table(
        slide,
        0.75,
        1.35,
        11.8,
        2.45,
        ["Run", "Samples", "Median", "P10", "P90", "Min"],
        [[row["stage"].replace("Final ", ""), row["sample_count"], row["median_compute_ms"], row["p10_compute_ms"], row["p90_compute_ms"], row["min_compute_ms"]] for row in final_rows],
        [2.0, 1.5, 1.8, 1.8, 1.8, 1.8],
        10.3,
    )
    metric_box(slide, 0.95, 4.55, 2.8, 1.1, "~0.99 ms", "Median band across repeats", "adopted")
    metric_box(slide, 4.25, 4.55, 2.8, 1.1, "~1.00 ms", "P90 band across repeats", "evidence")
    textbox(slide, 7.55, 4.75, 4.6, 0.8, "The confirmation set keeps median within a 0.0014 ms band and P90 near 1.00 ms across four runs.", size=14, color=C.INK, bold=True)
    page += 1

    slide = blank_slide(prs)
    header(slide, "Next experiment: reduce global spill", page)
    metric_box(slide, 0.95, 1.45, 3.05, 1.25, "Proven", "Shared queue atomics drove the measured bottleneck.", "evidence")
    metric_box(slide, 4.35, 1.45, 3.05, 1.25, "Kept", "Q1/Q2 sharding plus Q1 lane-pop stays in main.", "adopted")
    metric_box(slide, 7.75, 1.45, 3.05, 1.25, "Dropped", "Q2 request batching fails the default-suite test.", "rejected")
    rect(slide, 0.95, 3.65, 11.1, 1.4, fill=C.WHITE, line=C.LINE)
    textbox(slide, 1.25, 3.95, 10.5, 0.55, "The current architecture reaches about 0.99 ms median. The next useful experiment should reduce trips to global queues before revisiting batch claims.", size=15.5, color=C.INK, bold=True)
    page += 1

    slide = blank_slide(prs)
    header(slide, "Wave-local retention should measure global spill reduction", page)
    box_label(slide, 0.85, 1.6, 3.0, 1.2, "Current baseline", "Global sharded queues\nQ1/Q2 = 256/256", "adopted")
    arrow(slide, 4.05, 2.2, 4.85, 2.2, C.GRAY)
    box_label(slide, 5.0, 1.6, 3.0, 1.2, "Prototype", "Wave-local retention\nspill fallback", "evidence")
    arrow(slide, 8.2, 2.2, 9.0, 2.2, C.GRAY)
    box_label(slide, 9.15, 1.6, 3.0, 1.2, "Compare", "global atomic count\nspill ratio\nP90/tail stability", "neutral")
    multi_text(slide, 0.95, 3.75, 5.1, 1.2, ["Use the 256-shard main path as the baseline.", "Measure global atomic count, wave-local handoff count, spill ratio, and clean duration."], size=13, bullet=True)
    multi_text(slide, 6.75, 3.75, 5.1, 1.2, ["Bring request batching back after a queue-protocol change.", "Keep timing and atomic-counter diagnosis split in the next run."], size=13, bullet=True)
    page += 1

    slide = blank_slide(prs)
    header(slide, "Appendix: code surfaces behind each decision", page)
    snippet1 = """uint shard = mixHash(h) % Q1_QUEUE_SHARDS;
if (atomicCompSwap(control.q1[shard].count, cur, cur + 1u) == cur) {
    uint localSlot = atomicAdd(control.q1[shard].tail, 1u) % capacity;
    atomicExchange(queue1.tasks[slot].payload[5], 1u);
}"""
    snippet2 = """uint worker = wgId * local_size_x + localId;
uint homeShard =
    (ENABLE_Q1_LANE_POP != 0u)
        ? (worker % Q1_QUEUE_SHARDS)
        : (wgId % Q1_QUEUE_SHARDS);"""
    snippet3 = """uint claim = min(available, Q2_DEQUEUE_BATCH_LIMIT);
if (atomicCompSwap(control.q2[shard].count, cur, cur - claim) == cur) {
    baseHead = atomicAdd(control.q2[shard].head, claim);
}"""
    snippets = [
        ("SHARD PUSH", "Adopted: shard selection plus ready publish", snippet1, "adopted"),
        ("Q1 LANE-POP", "Adopted: lane-based Q1 scan start", snippet2, "adopted"),
        ("Q2 BATCH CAS", "Rejected: subgroup batch claim", snippet3, "rejected"),
    ]
    for i, (tag, title, code, kind) in enumerate(snippets):
        x = 0.7 + i * 4.15
        rect(slide, x, 1.35, 3.75, 4.5, fill=C.WHITE, line=C.LINE)
        semantic_tag(slide, x + 0.18, 1.58, tag, kind, w=1.55)
        textbox(slide, x + 0.18, 2.03, 3.35, 0.4, title, size=14, color=SEMANTIC[kind], bold=True)
        box = slide.shapes.add_textbox(Inches(x + 0.18), Inches(2.55), Inches(3.4), Inches(2.9))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        p = tf.paragraphs[0]
        p.text = code
        p.font.name = "Consolas"
        p.font.size = Pt(8.8)
        p.font.color.rgb = rgb(C.INK)
    page += 1

    slide = blank_slide(prs)
    header(slide, "Appendix: numeric source table and provenance", page)
    add_table(
        slide,
        0.65,
        1.35,
        12.0,
        3.0,
        ["Artifact", "Path"],
        [
            ["Normalized measurements", rel(CSV_OUT)],
            ["Measurement notes", rel(NOTES_OUT)],
            ["Main backbone", "reports/2026-05/metrics/q2_shards_20260510/research_backbone_20260510.md"],
            ["Request batching summary", "reports/2026-05/metrics/request_batching_validation_20260525/request_batching_validation_summary.csv"],
        ],
        [3.2, 8.8],
        9.3,
    )
    textbox(slide, 0.85, 5.05, 11.3, 0.65, "Use the normalized table for headline numbers; use the notes file for warmup rules and source-file mapping.", size=14.5, color=C.BLUE, bold=True)

    DECK_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(PPTX_OUT)


def main() -> None:
    rows = build_measurements()
    write_measurement_files(rows)
    assets = create_assets(rows)
    build_deck(rows, assets)
    print(f"Wrote {CSV_OUT}")
    print(f"Wrote {NOTES_OUT}")
    print(f"Wrote {PPTX_OUT}")


if __name__ == "__main__":
    main()
